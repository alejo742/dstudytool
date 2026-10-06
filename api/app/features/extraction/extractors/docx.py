from __future__ import annotations

from pathlib import Path

from app.features.extraction.extractors.base import ExtractedChunk, ExtractionResult


class DocxExtractor:
    def can_handle(self, mime: str, path: Path) -> bool:
        return (
            mime
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or path.suffix.lower() == ".docx"
        )

    def extract(self, path: Path, filename: str) -> ExtractionResult:
        from docx import Document

        doc = Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
        return ExtractionResult(
            chunks=[ExtractedChunk(text=text, source_label=filename)] if text else [],
            pages=None,
        )
